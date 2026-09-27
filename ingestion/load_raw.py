"""Ingesta EL: descarga Inside Airbnb (Barcelona) y las estadísticas de alquiler
de Barcelona Dades, y las carga sin transformar en la base `raw` de ClickHouse.

Uso:
    python ingestion/load_raw.py [--only airbnb|rent|hut|ine]

Idempotente: recargar un snapshot ya cargado reemplaza sus filas (DROP PARTITION
+ INSERT) sin duplicar.
"""

from __future__ import annotations

import argparse
import csv
import gc
import gzip
import io
import json
import logging
import os
import re
import zipfile
from collections import Counter
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Iterable

import clickhouse_connect
import requests

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
log = logging.getLogger("load_raw")

ROOT = Path(__file__).resolve().parent.parent
LANDING = ROOT / "data" / "landing"

# ponytail: Inside Airbnb no publica un índice de snapshots archivados por API
# documentada (la lista sale de una static query interna de su Gatsby site).
# Snapshots verificados a mano el 2026-09-23; añade fechas nuevas aquí cuando
# publiquen otro trimestre.
AIRBNB_SNAPSHOTS = ["2026-06-24", "2026-03-21", "2025-12-14", "2025-09-14"]
AIRBNB_DATA_ROOT = "https://data.insideairbnb.com/spain/catalonia/barcelona"

BCN_DADES_BASE = "https://portaldades.ajuntament.barcelona.cat"
BCN_RENT_STATS = {
    "b37xv8wcjh": "rent_price_monthly",  # €/mes
    "5ibudgqbrb": "rent_price_per_m2",  # €/m²
}

# Registro municipal de viviendas de uso turístico (HUT), un fichero por trimestre.
BCN_HUT_PACKAGE = (
    "https://opendata-ajuntament.barcelona.cat/data/api/3/action/package_show"
    "?id=habitatges-us-turistic"
)

# INE, estadística experimental de viviendas turísticas: un Excel por periodo con
# resultados por distrito censal (hoja 2) de toda España.
INE_VIV_TURISTICA = "https://www.ine.es/experimental/viv_turistica/"
_INE_MESES = {"FEB": 2, "MAY": 5, "AGO": 8, "NOV": 11}

CLICKHOUSE_DB = os.environ.get("CLICKHOUSE_DB", "raw")


def get_client():
    return clickhouse_connect.get_client(
        host=os.environ.get("CLICKHOUSE_HOST", "localhost"),
        port=int(os.environ.get("CLICKHOUSE_PORT", "8123")),
        username=os.environ.get("CLICKHOUSE_USER", "default"),
        password=os.environ.get("CLICKHOUSE_PASSWORD", ""),
        database=CLICKHOUSE_DB,
    )


def _create_raw_table(
    client, table: str, columns: list[str], partition_cols: list[str]
) -> None:
    cols_sql = ", ".join(f"`{c}` Nullable(String)" for c in columns)
    # las columnas son Nullable(String); ClickHouse no admite claves de partición nullable
    partition_expr = ", ".join(
        c if c == "_snapshot_date" else f"ifNull(`{c}`, '')" for c in partition_cols
    )
    client.command(
        f"""
        CREATE TABLE IF NOT EXISTS `{CLICKHOUSE_DB}`.`{table}` (
            {cols_sql},
            `_snapshot_date` Date,
            `_loaded_at` DateTime
        )
        ENGINE = MergeTree
        PARTITION BY ({partition_expr})
        ORDER BY tuple()
        """
    )
    # esquema fuente puede variar entre snapshots (p.ej. Inside Airbnb añade columnas nuevas)
    existing = {
        row[0]
        for row in client.query(
            f"DESCRIBE TABLE `{CLICKHOUSE_DB}`.`{table}`"
        ).result_rows
    }
    for c in columns:
        if c not in existing:
            client.command(
                f"ALTER TABLE `{CLICKHOUSE_DB}`.`{table}` ADD COLUMN IF NOT EXISTS `{c}` Nullable(String)"
            )


def _replace_partition(
    client,
    table: str,
    snapshot_date: str,
    columns: list[str],
    rows: Iterable[tuple],
    extra_partition: dict[str, str] | None = None,
) -> int:
    """Carga idempotente: tira la partición (snapshot + claves extra, p.ej. stat_id) e inserta de nuevo."""
    extra_partition = extra_partition or {}
    partition_cols = list(extra_partition.keys()) + ["_snapshot_date"]
    partition_values = list(extra_partition.values()) + [snapshot_date]

    _create_raw_table(client, table, columns, partition_cols)
    partition_literal = "(" + ", ".join(f"'{v}'" for v in partition_values) + ")"
    client.command(
        f"ALTER TABLE `{CLICKHOUSE_DB}`.`{table}` DROP PARTITION {partition_literal}"
    )

    snapshot_date_val = date.fromisoformat(snapshot_date)
    loaded_at = datetime.now(timezone.utc)
    all_cols = columns + ["_snapshot_date", "_loaded_at"]
    count = 0
    batch = []
    batch_size = 50_000
    for row in rows:
        batch.append((*row, snapshot_date_val, loaded_at))
        count += 1
        if len(batch) >= batch_size:
            client.insert(table, batch, column_names=all_cols)
            batch = []
            # cada insert deja ciclos de referencias (contexto de clickhouse_connect)
            # que el GC automático tarda en liberar: sin esto un calendario de ~7M
            # filas acumula ~1,5 GB y 4 snapshots seguidos agotan la memoria.
            gc.collect()
    if batch:
        client.insert(table, batch, column_names=all_cols)
    return count


def _download(url: str, dest: Path) -> Path:
    dest.parent.mkdir(parents=True, exist_ok=True)
    if dest.exists():
        log.info("ya en caché: %s", dest)
        return dest
    log.info("descargando %s", url)
    resp = requests.get(url, timeout=60)
    resp.raise_for_status()
    dest.write_bytes(resp.content)
    return dest


def _read_csv_gz(path: Path) -> tuple[list[str], Iterable[list[str]]]:
    f = gzip.open(path, "rt", encoding="utf-8", newline="")
    reader = csv.reader(f)
    header = next(reader)
    return header, reader


def load_airbnb_snapshot(client, snapshot_date: str) -> None:
    landing_dir = LANDING / "airbnb" / snapshot_date

    listings_gz = _download(
        f"{AIRBNB_DATA_ROOT}/{snapshot_date}/data/listings.csv.gz",
        landing_dir / "listings.csv.gz",
    )
    header, rows = _read_csv_gz(listings_gz)
    n = _replace_partition(client, "airbnb_listings", snapshot_date, header, rows)
    log.info("airbnb_listings %s: %d filas", snapshot_date, n)

    calendar_gz = _download(
        f"{AIRBNB_DATA_ROOT}/{snapshot_date}/data/calendar.csv.gz",
        landing_dir / "calendar.csv.gz",
    )
    header, rows = _read_csv_gz(calendar_gz)
    n = _replace_partition(client, "airbnb_calendar", snapshot_date, header, rows)
    log.info("airbnb_calendar %s: %d filas", snapshot_date, n)

    geojson_path = _download(
        f"{AIRBNB_DATA_ROOT}/{snapshot_date}/visualisations/neighbourhoods.geojson",
        landing_dir / "neighbourhoods.geojson",
    )
    geojson = json.loads(geojson_path.read_text(encoding="utf-8"))
    columns = [
        "neighbourhood",
        "neighbourhood_group",
        "geometry_type",
        "geometry_geojson",
    ]
    rows = [
        (
            f["properties"].get("neighbourhood"),
            f["properties"].get("neighbourhood_group"),
            f["geometry"]["type"],
            json.dumps(f["geometry"]),
        )
        for f in geojson["features"]
    ]
    n = _replace_partition(
        client, "airbnb_neighbourhoods", snapshot_date, columns, rows
    )
    log.info("airbnb_neighbourhoods %s: %d filas", snapshot_date, n)


def load_airbnb(client) -> None:
    for snapshot_date in AIRBNB_SNAPSHOTS:
        load_airbnb_snapshot(client, snapshot_date)


def _bcn_dades_client_id() -> str:
    resp = requests.get(
        f"{BCN_DADES_BASE}/portal/config/apimanagerconsumercredentials.json", timeout=30
    )
    resp.raise_for_status()
    return resp.json()["clientId"]


def _validate_rent_contract(payload: dict) -> None:
    """Falla claro si la API no documentada de Barcelona Dades cambia de forma."""
    if "dims" not in payload or "values" not in payload:
        raise ValueError("Contrato roto: la respuesta ya no tiene 'dims'/'values'")
    if not isinstance(payload["dims"], list) or len(payload["dims"]) < 2:
        raise ValueError("Contrato roto: 'dims' ya no tiene [tiempo, territorio]")
    if payload["dims"][0].get("id") != "dim0" or payload["dims"][1].get("id") != "dim1":
        raise ValueError("Contrato roto: los ids de 'dims' ya no son dim0/dim1")
    for v in payload["values"][:5]:
        if not {"dim0", "dim1", "value"} <= v.keys():
            raise ValueError(
                "Contrato roto: una fila de 'values' no tiene dim0/dim1/value"
            )


def _flatten_dim_tree(node: dict, index: dict) -> None:
    """Recorre el árbol de dims[N] y guarda id -> (type, label) para cada hoja/nodo."""
    if "id" in node:
        index[node["id"]] = (node.get("type"), node.get("label"))
    for child in node.get("values", []):
        _flatten_dim_tree(child, index)


def load_rent_stat(client, stat_id: str, snapshot_date: str) -> None:
    client_id = _bcn_dades_client_id()
    resp = requests.get(
        f"{BCN_DADES_BASE}/services/backend/rest/statistic",
        params={"id": stat_id, "language": "ca"},
        headers={"X-IBM-Client-Id": client_id},
        timeout=60,
    )
    resp.raise_for_status()
    payload = resp.json()
    _validate_rent_contract(payload)

    raw_path = LANDING / "bcn_rent" / f"{stat_id}_{snapshot_date}.json"
    raw_path.parent.mkdir(parents=True, exist_ok=True)
    raw_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")

    time_index: dict = {}
    for node in payload["dims"][0].get("values", []):
        _flatten_dim_tree(node, time_index)
    territory_index: dict = {}
    for node in payload["dims"][1].get("values", []):
        _flatten_dim_tree(node, territory_index)

    columns = [
        "stat_id",
        "period_id",
        "period_type",
        "territory_id",
        "territory_type",
        "territory_label",
        "value",
    ]
    rows = []
    for v in payload["values"]:
        period_type, _ = time_index.get(v["dim0"], (None, None))
        territory_type, territory_label = territory_index.get(v["dim1"], (None, None))
        rows.append(
            (
                stat_id,
                v["dim0"],
                period_type,
                v["dim1"],
                territory_type,
                territory_label,
                str(v["value"]) if v.get("value") is not None else None,
            )
        )

    n = _replace_partition(
        client,
        "bcn_rent",
        snapshot_date,
        columns,
        rows,
        extra_partition={"stat_id": stat_id},
    )
    log.info("bcn_rent[%s] %s: %d filas", stat_id, snapshot_date, n)


def load_rent(client) -> None:
    snapshot_date = datetime.now(timezone.utc).date().isoformat()
    for stat_id in BCN_RENT_STATS:
        load_rent_stat(client, stat_id, snapshot_date)


def _hut_snapshot_date(resource: dict) -> str:
    """'2026_1T_...' -> cierre del trimestre (2026-03-31); el fichero vigente
    (sin trimestre en el nombre) -> su fecha de publicación."""
    m = re.match(r"(\d{4})_(\d)T", resource["name"])
    if not m:
        return resource["last_modified"][:10]
    year, quarter = int(m[1]), int(m[2])
    next_quarter = date(year + quarter // 4, quarter % 4 * 3 + 1, 1)
    return (next_quarter - timedelta(days=1)).isoformat()


# nombres de columna que cambiaron entre trimestres -> nombre actual
_HUT_HEADER_ALIASES = {"DISTRICTE": "NOM_DISTRICTE", "BARRI": "NOM_BARRI"}
_HUT_ID = re.compile(r"^(HUTB-?\d+)?$")
# caja de Barcelona: una columna desplazada deja un número fuera de rango
_BCN_BOUNDS = {"LONGITUD_X": (2.0, 2.3), "LATITUD_Y": (41.3, 41.5)}


def _coord_ok(value: str, col: str) -> bool:
    if not value:
        return True
    try:
        lo, hi = _BCN_BOUNDS[col]
        return lo <= float(value.replace(",", ".")) <= hi
    except ValueError:
        return False


def _hut_row_ok(header: list[str], row: list[str]) -> bool:
    if len(row) != len(header):
        return False
    cell = dict(zip(header, row))
    return bool(_HUT_ID.match(cell.get("NUMERO_REGISTRE_GENERALITAT", ""))) and all(
        _coord_ok(cell.get(c, ""), c) for c in _BCN_BOUNDS
    )


def _repair_hut_header(header: list[str], lines: list[str], dialect) -> list[str]:
    """Cabeceras que no describen las filas (las filas son correctas):
    - 2018T2: 'LONGITUD_X -LATITUD_Y' en una sola columna.
    - 2020T3: faltan NUMERO_REGISTRE_GENERALITAT y NUMERO_PLACES, las filas sí los traen.
    Se decide por la longitud mayoritaria de las filas."""
    lengths = Counter(len(next(csv.reader([ln], dialect))) for ln in lines[:500])
    usual = lengths.most_common(1)[0][0]
    if "LONGITUD_X -LATITUD_Y" in header and usual == len(header) + 1:
        i = header.index("LONGITUD_X -LATITUD_Y")
        header = header[:i] + ["LONGITUD_X", "LATITUD_Y"] + header[i + 1 :]
    if (
        "NUMERO_REGISTRE_GENERALITAT" not in header
        and header[-2:] == ["LONGITUD_X", "LATITUD_Y"]
        and usual == len(header) + 2
    ):
        header = header[:-2] + [
            "NUMERO_REGISTRE_GENERALITAT",
            "NUMERO_PLACES",
            "LONGITUD_X",
            "LATITUD_Y",
        ]
    return header


def _repair_hut_row(header: list[str], line: str, row: list[str]) -> list[str] | None:
    """Repara las filas mal alineadas del histórico. Cada regla sale de un caso
    real (ver progress.md); si ninguna deja una fila válida, se descarta."""
    candidates = [
        # líneas enteras con ';' y coma decimal dentro de un fichero con ',' (2021T3)
        line.split(";"),
        # ',' y ';' mezclados en la misma línea (2022T1-T3)
        re.split(r"[,;]", line),
    ]
    # comas sin comillas en el nombre de barrio ("Sant Pere, Santa Caterina...")
    extra = len(row) - len(header)
    if extra > 0 and "NOM_BARRI" in header:
        i = header.index("NOM_BARRI")
        candidates.append(
            row[:i] + [",".join(row[i : i + extra + 1])] + row[i + extra + 1 :]
        )
    for cand in candidates:
        cand = [c.strip() for c in cand]
        if _hut_row_ok(header, cand):
            return cand
    return None


def _read_hut_csv(path: Path) -> tuple[list[str], list[list[str]], int, int]:
    """El histórico del Ajuntament no es homogéneo: .csv o .zip, utf-8 o latin-1,
    separador ',' o ';', y filas mal alineadas. Devuelve cabecera, filas válidas,
    nº de filas reparadas y nº de descartadas."""
    raw = path.read_bytes()
    if raw[:2] == b"PK":
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            raw = z.read(z.namelist()[0])
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")
    lines = [line for line in text.splitlines() if line.strip()]
    dialect = csv.Sniffer().sniff(lines[0], delimiters=",;")
    header = next(csv.reader([lines[0]], dialect))
    # cabecera sin salto de línea antes de la primera fila (2023T1, 2023T3):
    # "...,LATITUD_Y01-2013-0753,1,CIUTAT VELLA,..."
    glued = next(
        (
            i
            for i, c in enumerate(header)
            if c.startswith("LATITUD_Y") and c != "LATITUD_Y"
        ),
        None,
    )
    if glued is not None:
        first = [header[glued][len("LATITUD_Y") :]] + header[glued + 1 :]
        header = header[:glued] + ["LATITUD_Y"]
        lines[0] = dialect.delimiter.join(first)
    else:
        lines = lines[1:]
    header = [_HUT_HEADER_ALIASES.get(c, c) for c in header]
    header = _repair_hut_header(header, lines, dialect)

    rows, repaired, dropped = [], 0, 0
    for line in lines:
        row = next(csv.reader([line], dialect))
        if len(row) == len(header):
            rows.append(row)
            continue
        fixed = _repair_hut_row(header, line, row)
        if fixed is None:
            dropped += 1
        else:
            rows.append(fixed)
            repaired += 1
    return header, rows, repaired, dropped


def load_hut_registry(client) -> None:
    resp = requests.get(BCN_HUT_PACKAGE, timeout=30)
    resp.raise_for_status()
    for resource in resp.json()["result"]["resources"]:
        snapshot_date = _hut_snapshot_date(resource)
        # la fecha va en el nombre: el fichero vigente se republica con la misma URL
        path = _download(
            resource["url"], LANDING / "bcn_hut" / f"{snapshot_date}_{resource['name']}"
        )
        header, rows, repaired, dropped = _read_hut_csv(path)
        # antes de 2020T4 (y en 2024T2) no se publica el nº HUTB: esos trimestres
        # no sirven para cruzar con los anuncios, pero sí para contar licencias.
        if repaired or dropped:
            log.warning(
                "bcn_hut %s: %d filas reparadas, %d descartadas",
                resource["name"],
                repaired,
                dropped,
            )
        n = _replace_partition(client, "bcn_hut_registry", snapshot_date, header, rows)
        log.info("bcn_hut_registry %s: %d filas", snapshot_date, n)


def _read_ine_districts(path: Path) -> tuple[list[str], list[tuple]]:
    """Hoja de distritos del Excel del INE. Las cabeceras cambian entre periodos
    (CODIGO/CUDIS, Prov/PROV...): se normalizan a minúsculas con '_'."""
    import openpyxl  # solo lo necesita esta fuente

    wb = openpyxl.load_workbook(path, read_only=True)
    rows = wb[wb.sheetnames[1]].iter_rows(values_only=True)
    header = [str(c).strip().lower().replace(" ", "_") for c in next(rows)]
    header = ["codigo" if c == "cudis" else c for c in header]
    body = [tuple(None if v is None else str(v) for v in r) for r in rows if r and r[0]]
    wb.close()
    return header, body


def load_ine_tourist_dwellings(client) -> None:
    resp = requests.get(
        f"{INE_VIV_TURISTICA}exp_viv_turistica_descarga.htm", timeout=30
    )
    resp.raise_for_status()
    for name in sorted(
        set(re.findall(r"exp_viv_turistica_tabla5_\w+\.xlsx", resp.text))
    ):
        month, year = re.search(r"_([A-Z]{3})(\d{4})\.xlsx", name).groups()
        snapshot_date = date(int(year), _INE_MESES[month], 1).isoformat()
        path = _download(f"{INE_VIV_TURISTICA}{name}?nocab=1", LANDING / "ine" / name)
        header, rows = _read_ine_districts(path)
        n = _replace_partition(
            client, "ine_tourist_dwellings", snapshot_date, header, rows
        )
        log.info("ine_tourist_dwellings %s: %d filas", snapshot_date, n)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--only", choices=["airbnb", "rent", "hut", "ine"], default=None
    )
    args = parser.parse_args()

    client = get_client()
    client.command(f"CREATE DATABASE IF NOT EXISTS `{CLICKHOUSE_DB}`")

    if args.only in (None, "airbnb"):
        load_airbnb(client)
    if args.only in (None, "rent"):
        load_rent(client)
    if args.only in (None, "hut"):
        load_hut_registry(client)
    if args.only in (None, "ine"):
        load_ine_tourist_dwellings(client)


if __name__ == "__main__":
    main()
